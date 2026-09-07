import os
import shutil
from datetime import datetime
from openpyxl import load_workbook


# ==========================================
# 設定
# ==========================================

# Excelファイル
excel_file = "礼拝奉仕ローテーション.xlsx"

# ファイルを探すフォルダ
source_folder = "psalm_jpg"

# コピー先フォルダ
destination_folder = "psalm"


# ==========================================
# 詩篇から「148A」のような名前を抽出
# ==========================================

def extract_psalm_name(text):

    if not text:
        return ""

    text = str(text).strip()

    # 「詩篇」または「詩編」を削除
    if text.startswith("詩篇"):
        text = text[2:]

    elif text.startswith("詩編"):
        text = text[2:]

    # "("以降を削除
    #
    # 例：
    # 148A(1,5,7)
    # ↓
    # 148A

    text = text.split("(")[0]

    return text.strip()


# ==========================================
# 聖歌から「390」のような番号を抽出
# ==========================================

def extract_hymn_name(text):

    if not text:
        return ""

    text = str(text).strip()

    # 「聖歌」を削除
    #
    # 例：
    # 聖歌390(1,3,4)
    # ↓
    # 390(1,3,4)

    if text.startswith("聖歌"):
        text = text[2:]

    # "("以降を削除
    #
    # 390(1,3,4)
    # ↓
    # 390

    text = text.split("(")[0]

    return text.strip()


# ==========================================
# コピー先フォルダを作成
# ==========================================

os.makedirs(
    destination_folder,
    exist_ok=True
)


# ==========================================
# Excelを読み込む
# ==========================================

try:

    wb = load_workbook(
        excel_file,
        data_only=True
    )

    ws = wb["2025.11.30-2026.11.22"]

except FileNotFoundError:

    print(
        f"Excelファイルが見つかりません：{excel_file}"
    )

    input("Enterキーを押して終了...")
    exit()

except KeyError:

    print(
        "指定したシートが見つかりません。"
    )

    input("Enterキーを押して終了...")
    exit()


# ==========================================
# Excelを1行ずつ処理
# ==========================================

for row in ws.iter_rows(
    min_row=2,
    values_only=True
):

    # ======================================
    # B列 → 日付
    # ======================================

    date_value = row[1]


    # ======================================
    # Z列 → 詩篇・聖歌
    # ======================================

    cell_value = row[25]


    # ======================================
    # 空欄ならスキップ
    # ======================================

    if not date_value or not cell_value:
        continue


    # ======================================
    # 日付を文字列に変換
    # ======================================

    if isinstance(
        date_value,
        datetime
    ):

        date_string = date_value.strftime(
            "%Y-%m-%d"
        )

    else:

        date_string = str(
            date_value
        ).strip()


    # ======================================
    # 詩篇か聖歌か判定
    # ======================================

    cell_text = str(
        cell_value
    ).strip()


    # ======================================
    # 詩篇の場合
    # ======================================

    if (
        cell_text.startswith("詩篇")
        or cell_text.startswith("詩編")
    ):

        # ----------------------------------
        # 詩篇番号を抽出
        # ----------------------------------

        extracted_name = extract_psalm_name(
            cell_text
        )


        if not extracted_name:

            print(
                f"詩篇番号を抽出できませんでした：{cell_text}"
            )

            continue


        # ----------------------------------
        # 検索するファイル名
        # ----------------------------------

        search_name = (
            "psalm"
            + extracted_name
        )


        # ----------------------------------
        # コピー後のファイル名
        # ----------------------------------

        final_name = (
            "psalm"
            + date_string
            + ".jpg"
        )


        file_type = "詩篇"


    # ======================================
    # 聖歌の場合
    # ======================================

    elif cell_text.startswith("聖歌"):

        # ----------------------------------
        # 聖歌番号を抽出
        # ----------------------------------

        extracted_name = extract_hymn_name(
            cell_text
        )


        if not extracted_name:

            print(
                f"聖歌番号を抽出できませんでした：{cell_text}"
            )

            continue


        # ----------------------------------
        # 検索するファイル名
        # ----------------------------------

        search_name = (
            "hymn"
            + extracted_name
        )


        # ----------------------------------
        # コピー後のファイル名
        # ----------------------------------

        final_name = (
            "hymn"
            + date_string
            + ".jpg"
        )


        file_type = "聖歌"


    # ======================================
    # その他の場合
    # ======================================

    else:

        print(
            f"【種類不明】{cell_text}"
        )

        continue


    # ======================================
    # 処理内容を表示
    # ======================================

    print("----------------------------------")

    print(
        f"日付          ：{date_string}"
    )

    print(
        f"元のセル      ：{cell_text}"
    )

    print(
        f"種類          ：{file_type}"
    )

    print(
        f"抽出した番号  ：{extracted_name}"
    )

    print(
        f"検索する名前  ：{search_name}"
    )

    print(
        f"新しい名前    ：{final_name}"
    )


    # ======================================
    # ファイルを探す
    # ======================================

    found = False


    for root, dirs, files in os.walk(
        source_folder
    ):

        for filename in files:

            # 拡張子を除いたファイル名
            name_without_extension = os.path.splitext(
                filename
            )[0]


            # ==================================
            # ファイル名が一致
            # ==================================

            if name_without_extension == search_name:

                source_file = os.path.join(
                    root,
                    filename
                )


                # ==================================
                # コピー先
                # ==================================

                destination_file = os.path.join(
                    destination_folder,
                    final_name
                )


                # ==================================
                # コピー
                # ==================================

                shutil.copy2(
                    source_file,
                    destination_file
                )


                print(
                    "コピー完了！"
                )

                print(
                    f"  元ファイル ：{source_file}"
                )

                print(
                    f"  新ファイル ：{destination_file}"
                )


                found = True

                break


        # ファイルが見つかったら
        # サブフォルダ検索を終了

        if found:
            break


    # ======================================
    # ファイルが見つからなかった場合
    # ======================================

    if not found:

        print(
            f"【見つかりません】{search_name}"
        )


# ==========================================
# 終了
# ==========================================

print()
print("==========================================")
print("すべての処理が完了しました。")
print("==========================================")

input("Enterキーを押して終了...")